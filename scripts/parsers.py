"""
Readers for the college's files. Each takes a file and returns plain data;
none of them writes anything. update.py decides what to do with the result.

Three kinds of file are understood, whatever the programme:

  weekly schedule (PDF)   a ruled grid, one column per time slot
  course allocation (PDF) a table of courses with their abbreviations
  student list (xls, xlsx or PDF) a table with a "Roll No" column

Schedules come in three layouts, told apart by the columns left of the time
slots:

  AREA column     rows are subject areas; a cell names course and section,
                  e.g. "CWB-C(1) PJ{C -201}". A student is matched by the
                  course and section they registered for.   (mode "course")
  Sec column      rows are sections with their room; a cell names the course,
                  e.g. "SERM-(1)-SPP" or "HRM(1)FH". A student is matched by
                  their section.                            (mode "section")
  Track column    rows are elective tracks, e.g. "VB-A(13)-MZ" in row AQF.
                  A student is matched by their track.      (mode "track")
"""
import datetime as dt
import re

MONTHS = {m: i for i, m in enumerate("jan feb mar apr may jun jul aug sep oct nov dec".split(), 1)}
WEEKDAYS = {d: i for i, d in enumerate("mon tue wed thu fri sat sun".split())}
ROMAN = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6}


class ParseError(Exception):
    pass


def clean(s):
    return re.sub(r"\s+", " ", str(s if s is not None else "")).strip()


def norm_code(s):
    """Course abbreviations compare equal whatever their punctuation: DSL-1 = DSL1."""
    return re.sub(r"[^A-Z0-9&]", "", str(s).upper())


def norm_name(s):
    return re.sub(r"[^a-z0-9]+", " ", str(s).lower()).strip()


def clean_roll(v):
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return re.sub(r"[^A-Z0-9]", "", str(v).upper())


def clean_person(v):
    name = clean(v)
    name = re.sub(r"(^| )\.( |$)", " ", name)  # "Dhruv ." -> "Dhruv"
    return clean(name)


# ----------------------------------------------------------------- file access

def pdf_tables(path):
    """[(page_text, [table_rows])] for a PDF; cells keep None where a merged
    cell covers them, which is how spans are recognised later."""
    import pymupdf
    doc = pymupdf.open(str(path))
    out = []
    for page in doc:
        out.append((page.get_text(), [t.extract() for t in page.find_tables().tables]))
    return out


def sheet_rows(path):
    """Every row of every sheet in a workbook, one sheet after another. A class
    list is often one sheet per section, so stopping at the first sheet would
    silently drop the rest. Each sheet's own header row is recognised again by
    whoever reads the rows."""
    suffix = path.suffix.lower()
    rows = []
    if suffix == ".xls":
        import xlrd
        for ws in xlrd.open_workbook(str(path)).sheets():
            rows += [[ws.cell_value(r, c) for c in range(ws.ncols)] for r in range(ws.nrows)]
        return rows
    import openpyxl
    for ws in openpyxl.load_workbook(str(path), data_only=True).worksheets:
        rows += [["" if c is None else c for c in row] for row in ws.iter_rows(values_only=True)]
    return rows


# ------------------------------------------------------------ what is this file

def describe(path):
    """Work out what a file is from its contents: {'kind', 'batch', 'group', 'term'}.
    kind is 'weekly', 'courses', 'students' or None."""
    info = {"kind": None, "batch": None, "group": None, "term": None}
    suffix = path.suffix.lower()
    if suffix in (".xls", ".xlsx"):
        rows = sheet_rows(path)
        text = " ".join(clean(c) for row in rows[:6] for c in row) + " " + path.name
        rolls = [clean_roll(c) for row in rows for c in row if re.fullmatch(r"\d{9}", clean_roll(c))]
        if rolls:
            info["kind"] = "students"
            info["batch"] = "20%02d" % (int(rolls[0][:2]) + 2)
            info["_code"] = rolls[0][2:6]
    elif suffix == ".pdf":
        pages = pdf_tables(path)
        text = "\n".join(p[0] for p in pages) + " " + path.name
        rows = [r for _, tables in pages for t in tables for r in t]
        if any(sum(1 for c in r if parse_slot(clean(c))) >= 3 for r in rows):
            info["kind"] = "weekly"
        elif any(any(re.fullmatch(r"(roll|enrol+ment)\s*(no\.?|number)", clean(c).lower()) for c in r) for r in rows):
            info["kind"] = "students"
            rolls = [clean_roll(c) for r in rows for c in r if re.fullmatch(r"\d{9}", clean_roll(c))]
            if rolls:
                info["batch"] = "20%02d" % (int(rolls[0][:2]) + 2)
                info["_code"] = rolls[0][2:6]
        elif any(any("abb" in clean(c).lower() for c in r) or sum(1 for c in r if clean(c).lower() == "course") >= 2 for r in rows):
            info["kind"] = "courses"
    else:
        return info
    m = re.search(r"(20\d{2})\s*[-–]\s*(\d{2})\b", text)
    if m and not info["batch"]:
        info["batch"] = "20" + m.group(2)
    upper = text.upper()
    info["group"] = "dcp" if "DCP" in upper else "bfs" if "BFS" in upper else "core"
    m = re.search(r"TERM\s*[-–]?\s*(VI|IV|V|III|II|I|[1-6])\b", upper) or re.search(r"\bT\s*-\s*(VI|IV|V|III|II|I)\b", upper)
    if m:
        info["term"] = ROMAN.get(m.group(1)) or int(m.group(1))
    return info


# ------------------------------------------------------------------- schedules

def parse_slot(text):
    """'11:45 am -01:00 pm', '09:00 AM - 10:15 AM', '9:00 - 10:15 Hrs.' -> ('11:45', '13:00')."""
    t = text.lower().replace(".", "")
    m = re.fullmatch(r"\s*(\d{1,2}):(\d{2})\s*(am|pm)?\s*[-–]\s*(\d{1,2}):(\d{2})\s*(am|pm|hrs)?\s*", t)
    if not m:
        return None
    h1, m1, mer1, h2, m2, mer2 = m.groups()
    if mer2 in (None, "hrs") and not mer1:
        start, end = int(h1) * 60 + int(m1), int(h2) * 60 + int(m2)
    else:
        def to24(h, mi, mer):
            return int(h) % 12 * 60 + (720 if mer == "pm" else 0) + int(mi)
        end = to24(h2, m2, mer2)
        start = to24(h1, m1, mer1 or mer2)
        if start >= end and not mer1:
            start = to24(h1, m1, "am")
    if not 0 <= start < end <= 24 * 60:
        return None
    return ("%02d:%02d" % divmod(start, 60), "%02d:%02d" % divmod(end, 60))


def parse_day(text, years):
    """'Mon, Oct. 05, 2026', '05/10/2026 (Mon)', 'Mon, Oct 05' -> date;
    a bare weekday such as 'Thu' -> its index 0-6; anything else -> None.
    A date with no year takes the year (from those the file mentions) in which
    it falls on the weekday written next to it."""
    t = text.strip().lower()
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", t)
    if m:
        return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    m = re.search(r"([a-z]{3})[a-z]*\.?\s+(\d{1,2}),?\s+(\d{4})", t)
    if m and m.group(1) in MONTHS:
        return dt.date(int(m.group(3)), MONTHS[m.group(1)], int(m.group(2)))
    m = re.search(r"([a-z]{3})[a-z]*\.?,?\s+([a-z]{3})[a-z]*\.?\s+(\d{1,2})\b", t)
    if m and m.group(1) in WEEKDAYS and m.group(2) in MONTHS:
        for year in sorted(years) + list(range(2024, 2040)):
            d = dt.date(year, MONTHS[m.group(2)], int(m.group(3)))
            if d.weekday() == WEEKDAYS[m.group(1)]:
                return d
    if t[:3] in WEEKDAYS and len(t) <= 9:
        return WEEKDAYS[t[:3]]
    return None


def parse_class(text):
    """One timetable cell -> {'c', 'sec', 'grp', 'n', 'room'} or None.

        CWB-C(1) PJ{C -201}   course CWB, section C, session 1, room C-201
        SERM-(1)-SPP          course SERM, session 1
        RMBFS-I-(1)-SDG       course RMBFS-I, session 1
        VB-A(13)-MZ           course VB, section A, session 13
        HRM(1)FH              course HRM, session 1
        DTI-G-5(1)AT          course DTI, group 5, session 1
    """
    m = re.fullmatch(r"(.+?)\s*\(\s*(\d+)\s*\)\s*(.*)", text)
    if not m:
        return None
    head, num, tail = m.group(1).strip(), int(m.group(2)), m.group(3)
    sec = grp = None
    if head.endswith("-"):
        head = head[:-1].strip()
    else:
        g = re.fullmatch(r"(.+?)\s*-\s*G\s*-\s*(\d+)", head)
        s = re.fullmatch(r"(.+?)\s*-\s*([A-Z])", head)
        if g:
            head, grp = g.group(1), g.group(2)
        elif s:
            head, sec = s.group(1), s.group(2)
    if not re.fullmatch(r"[A-Z][A-Z0-9&]*(?:-[A-Z0-9]+)*", head):
        return None
    room = None
    r = re.search(r"\{([^{}]*)\}", tail)
    if r:
        room = re.sub(r"\s*-\s*", "-", clean(r.group(1)))
    return {"c": head, "sec": sec, "grp": grp, "n": num, "room": room}


def day_bands(path, years):
    """Which day each table row belongs to, worked out from where the day
    labels sit on the page: {(page index, table index): [day per row]}.

    A day label ("Tue, Oct 13, 2026") is one merged cell beside that day's
    rows, written sideways and centred in it. The table finder sometimes
    cuts that cell into pieces, so the label's text arrives as fragments
    ("13, 2026" on one row, "Tue, Oct" three rows down) or lands on the wrong
    row. The page itself still has each label as one line of text with a
    position, and a label centred in its cell fixes that cell's edges: the
    first day's band starts at the first data row and extends as far below
    the label's centre as it does above; the next band starts where that one
    ends; and so on. A row belongs to the band its own centre falls in.

    A table is left out (so the cell text is used instead) when it has no
    readable labels or the bands do not account for its rows.
    """
    import pymupdf
    out = {}
    doc = pymupdf.open(str(path))
    for p_index, page in enumerate(doc):
        lines = []
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                text = clean("".join(span["text"] for span in line["spans"]))
                if text:
                    lines.append((line["bbox"], text))
        for t_index, table in enumerate(page.find_tables().tables):
            rows = table.extract()
            boxes = [r.bbox for r in table.rows]
            if len(boxes) != len(rows):
                continue
            header = [i for i, r in enumerate(rows) if sum(1 for c in r if parse_slot(clean(c))) >= 3]
            first = header[-1] + 1 if header else 0
            if first >= len(rows):
                continue
            x0, top, x1, bottom = table.bbox[0], boxes[first][1], table.bbox[2], boxes[-1][3]
            labels = []
            for (lx0, ly0, lx1, ly1), text in lines:
                centre_x, centre_y = (lx0 + lx1) / 2, (ly0 + ly1) / 2
                if not (top <= centre_y <= bottom and x0 - 2 <= centre_x <= x0 + 0.25 * (x1 - x0)):
                    continue
                day = parse_day(text, years)
                if day is not None:
                    labels.append((centre_y, day))
            labels.sort(key=lambda l: l[0])
            if not labels:
                continue
            bands, edge = [], top
            for centre_y, day in labels:
                height = 2 * (centre_y - edge)
                if height <= 4:
                    bands = None
                    break
                bands.append((edge, edge + height, day))
                edge += height
            if not bands or abs(edge - bottom) > 14:
                continue
            days = [None] * first
            for box in boxes[first:]:
                centre = (box[1] + box[3]) / 2
                hit = [d for lo, hi, d in bands if lo <= centre < hi]
                days.append(hit[0] if hit else bands[-1][2])
            out[(p_index, t_index)] = days
    return out


def parse_weekly(path):
    """Read one weekly schedule PDF.

    Returns {'term', 'weekNumber', 'start', 'end', 'mode', 'rowLabel', 'rows',
    'slots', 'days', 'sessions', 'specials', 'warnings'}.
    Raises ParseError rather than return a week it isn't sure it read fully.
    """
    pages = pdf_tables(path)
    raw_text = "\n".join(p[0] for p in pages)
    years = {int(y) for y in re.findall(r"\b(20\d{2})\b", raw_text + " " + path.name)}
    raw_count = len(re.findall(r"\(\s*\d+\s*\)", raw_text))

    slots = slot_cols = None
    cols = {}                 # role -> column index: day, row, room, area
    mode = None
    sessions, specials, notes, warnings = [], [], {}, []
    row_keys = []
    day = None
    day_rows = []             # rows of the current day: (row_key, raw_cells), for merged cells
    day_specials = []

    def close_day():
        # a merged cell shows as None in the rows under it: those rows share the entry
        for sp, first_row in day_specials:
            col = slot_cols[sp["s"]]
            for i in range(first_row + 1, len(day_rows)):
                key, cells = day_rows[i]
                if col < len(cells) and cells[col] is None and key and key not in sp["rows"]:
                    sp["rows"].append(key)
                else:
                    break
        del day_rows[:], day_specials[:]

    bands = day_bands(path, years)
    for p_index, (_, tables) in enumerate(pages):
        for t_index, table in enumerate(tables):
            banded = bands.get((p_index, t_index))
            for r_index, raw in enumerate(table):
                cells = [clean(c) for c in raw]
                parsed = [parse_slot(c) for c in cells]
                if sum(1 for p in parsed if p) >= 3:
                    found_cols = [i for i, p in enumerate(parsed) if p]
                    found = [parsed[i] for i in found_cols]
                    if slots is not None and found != slots:
                        raise ParseError("time slots differ between pages: %s vs %s" % (slots, found))
                    slots, slot_cols = found, found_cols
                    above = [clean(c) for c in table[r_index - 1]] if r_index else []
                    cols = {}
                    for i in range(found_cols[0]):
                        label = (cells[i] + " " + (above[i] if i < len(above) else "")).lower()
                        if "day" in label:
                            cols.setdefault("day", i)
                        elif "track" in label:
                            cols["row"], mode = i, "track"
                        elif "sec" in label:
                            cols["row"], mode = i, "section"
                        elif "area" in label:
                            cols["area"], mode = i, "course"
                        elif "venue" in label or "classroom" in label or "room" in label:
                            cols["room"] = i
                    if "day" not in cols or mode is None:
                        raise ParseError("could not tell the layout from the header row: %s" % cells[:found_cols[0]])
                    continue
                if slots is None:
                    continue  # title rows above the first header

                day_text = cells[cols["day"]] if cols["day"] < len(cells) else ""
                if banded:
                    # the day comes from where the row sits on the page; a whole date
                    # written in the row itself must agree with it
                    placed = banded[r_index]
                    written = parse_day(day_text, years) if day_text else None
                    if isinstance(written, dt.date) and isinstance(placed, dt.date) and written != placed:
                        raise ParseError("a row labelled %s sits in the band of %s" % (written, placed))
                    if placed != day:
                        close_day()
                        day = placed
                elif day_text:
                    close_day()
                    day = parse_day(day_text, years)
                    if day is None:
                        raise ParseError("unreadable day label: %r" % day_text)
                if day is None:
                    continue
                row_key = cells[cols["row"]] if "row" in cols and cols["row"] < len(cells) else None
                row_room = cells[cols["room"]] if "room" in cols and cols["room"] < len(cells) else None
                if row_key and row_key not in row_keys:
                    row_keys.append(row_key)
                day_rows.append((row_key, raw))

                found_any = False
                for s_idx, col in enumerate(slot_cols):
                    text = cells[col] if col < len(cells) else ""
                    if not text:
                        continue
                    cls = parse_class(text)
                    if cls:
                        found_any = True
                        cls.update(d=day, s=s_idx, row=row_key)
                        cls["room"] = cls["room"] or re.sub(r"\s*-\s*", "-", row_room or "") or None
                        sessions.append(cls)
                    elif len(text) > 2:  # stray single characters are scan noise
                        found_any = True
                        span = 1
                        while (s_idx + span < len(slot_cols) and slot_cols[s_idx + span] < len(raw)
                               and raw[slot_cols[s_idx + span]] is None):
                            span += 1
                        sp = {"d": day, "s": s_idx, "span": span, "text": text, "rows": [row_key] if row_key else []}
                        specials.append(sp)
                        day_specials.append((sp, len(day_rows) - 1))
                if not found_any and mode == "course":
                    # e.g. "Thu | SSR Visits": a whole-day note written across the row
                    extra = " ".join(c for i, c in enumerate(cells) if i != cols["day"] and c
                                     and not re.fullmatch(r"[A-Z]{2,4}", c))
                    if extra:
                        notes[day] = (notes.get(day, "") + " " + extra).strip()
    close_day()

    if slots is None:
        raise ParseError("no time-slot header row found")
    if len(sessions) != raw_count:
        raise ParseError("the PDF text has %d class entries but %d were read from the grid" % (raw_count, len(sessions)))

    dated = sorted({x for x in [s["d"] for s in sessions] + [s["d"] for s in specials] + list(notes) if isinstance(x, dt.date)})
    if not dated:
        raise ParseError("no dated day labels found")
    monday = dated[0] - dt.timedelta(days=dated[0].weekday())
    if any((d - monday).days > 6 for d in dated):
        raise ParseError("days span more than one Monday-Sunday week: %s to %s" % (dated[0], dated[-1]))

    def day_index(d):
        return (d - monday).days if isinstance(d, dt.date) else d

    for item in sessions + specials:
        item["d"] = day_index(item["d"])
    notes = {day_index(k): v for k, v in notes.items()}
    sessions.sort(key=lambda x: (x["d"], x["s"], x["row"] or "", x["c"], x["sec"] or "", x["grp"] or ""))

    streams = {}
    for x in sessions:
        streams.setdefault((x["c"], x["sec"], x["grp"], x["row"]), []).append(x["n"])
    for key, nums in sorted(streams.items(), key=str):
        if sorted(nums) != list(range(min(nums), min(nums) + len(nums))):
            label = "-".join(str(k) for k in key if k)
            warnings.append("%s session numbers run %s, expected consecutive" % (label, nums))

    upper = (raw_text + " " + path.name).upper()
    m = re.search(r"TERM\s*[-–]?\s*(VI|IV|V|III|II|I|[1-6])\b", upper)
    term = (ROMAN.get(m.group(1)) or int(m.group(1))) if m else None
    w = (re.search(r"WEEK\s*[-–]?\s*(\d+)", upper) or re.search(r"\b(\d+)\s*(?:ST|ND|RD|TH)\s+WEEKLY", upper)
         or re.match(r"\s*(\d+)\s*-\s*WEEKLY", path.name.upper()))
    days = []
    for i in range(7):
        entry = {"date": (monday + dt.timedelta(days=i)).isoformat()}
        if i in notes:
            entry["note"] = notes[i]
        days.append(entry)
    return {
        "term": term,
        "weekNumber": int(w.group(1)) if w else None,
        "start": monday.isoformat(),
        "end": (monday + dt.timedelta(days=6)).isoformat(),
        "mode": mode,
        "rowLabel": {"section": "Section", "track": "Track", "course": None}[mode],
        "rows": row_keys,
        "slots": [list(s) for s in slots],
        "days": days,
        "sessions": sessions,
        "specials": specials,
        "warnings": warnings,
    }


# --------------------------------------------------------------------- courses

def parse_courses(path):
    """Course allocation sheet -> {abbreviation: {'name', 'area', 'credits', 'sections'}}."""
    courses, warnings = {}, []
    for _, tables in pdf_tables(path):
        for table in tables:
            idx = None
            last = {}
            for raw in table:
                cells = [clean(c) for c in raw]
                low = [c.lower() for c in cells]
                if idx is None:
                    abb = [i for i, c in enumerate(low) if "abb" in c and "fac" not in c]
                    plain = [i for i, c in enumerate(low) if c == "course"]
                    name = [i for i, c in enumerate(low) if c in ("course", "course name")]
                    if not name or not (abb or len(plain) >= 2):
                        continue
                    idx = {"name": name[0], "abb": abb[0] if abb else plain[1]}
                    for key, test in (("area", lambda c: c == "area"), ("credits", lambda c: c.startswith("cr")),
                                      ("sec", lambda c: c in ("sec", "sec."))):
                        hit = [i for i, c in enumerate(low) if test(c)]
                        if hit:
                            idx[key] = hit[0]
                    continue
                get = lambda key: cells[idx[key]] if key in idx and idx[key] < len(cells) else ""
                abb = get("abb").replace(" ", "")
                if not abb or not re.fullmatch(r"[A-Za-z][A-Za-z0-9&\-]*", abb):
                    continue
                name = get("name") or last.get("name", "")
                last = {"name": name}
                if not name:
                    continue
                entry = {"name": name}
                if get("area"):
                    entry["area"] = re.sub(r"\s+", "", get("area")) if len(get("area")) <= 5 else get("area")
                if re.fullmatch(r"\d+(\.\d+)?", get("credits")):
                    entry["credits"] = float(get("credits"))
                if re.fullmatch(r"\d+", get("sec")):
                    entry["sections"] = int(get("sec"))
                key = abb.upper()
                if key in courses and norm_name(courses[key]["name"]) != norm_name(name):
                    warnings.append("abbreviation %s is used for both %r and %r in %s"
                                    % (key, courses[key]["name"], name, path.name))
                    continue
                courses.setdefault(key, entry)
    if not courses:
        raise ParseError("no course table found")
    return courses, warnings


# -------------------------------------------------------------------- students

STUDENT_COLUMNS = {
    "roll no": "roll", "roll no.": "roll", "roll number": "roll", "roll": "roll",
    "enrolment no": "roll", "enrolment no.": "roll", "enrollment no": "roll", "enrollment no.": "roll",
    "enrolment number": "roll", "enrollment number": "roll",
    "student name": "name", "name": "name",
    "sec": "section", "sec.": "section", "section": "section",
    "major": "major", "minor": "minor", "track": "track", "total": "total",
    "elective": "elective", "elective course": "elective",
}


def match_course(name, courses):
    """The abbreviation of the course a written-out name refers to. Names in
    different files differ in small words ("Service Excellence in ..." vs
    "Service Excellence & ..."), so the best overlap of words wins, if it is
    clearly the best."""
    want = set(norm_name(name).replace("behavioral", "behavioural").split()) - {"and", "in", "of", "for", "the"}
    scored = []
    for abb, c in courses.items():
        have = set(norm_name(c["name"]).replace("behavioral", "behavioural").split()) - {"and", "in", "of", "for", "the"}
        if want and have:
            scored.append((len(want & have) / len(want | have), abb))
    scored.sort(reverse=True)
    if scored and scored[0][0] >= 0.6 and (len(scored) == 1 or scored[0][0] - scored[1][0] >= 0.2):
        return scored[0][1]
    return None


def parse_students(path, courses=None):
    """Student list -> ({roll: {'name', 'section', 'major', 'minor', 'track', 'courses'}}, notes).
    An "Elective Course" column becomes the student's elective list.
    Only these fields are ever read; emails, dates of birth and the like in a
    class profile are ignored. Columns headed with a course's full name (the
    Term 4-6 registration sheet) become that student's course -> section map,
    and so does a sheet with one row per student and course ("Course" holding
    the abbreviation, "SEC" that course's section)."""
    by_name = {norm_name(c["name"]): abb for abb, c in (courses or {}).items()}
    if path.suffix.lower() == ".pdf":
        rows = [r for _, tables in pdf_tables(path) for t in tables for r in t]
    else:
        rows = sheet_rows(path)
    students, problems, notes, conflicted = {}, [], [], set()
    by_code = {norm_code(abb): abb for abb in (courses or {})}
    columns, long_form = None, False
    for raw in rows:
        cells = [clean(c) if not isinstance(c, float) else c for c in raw]
        low = [clean(c).lower() for c in raw]
        if any(STUDENT_COLUMNS.get(c) == "roll" and c != "roll" for c in low):
            columns = {}
            # one row per student and course ("Roll No | Name | Course | SEC"):
            # the Course column holds an abbreviation and SEC is that course's section
            long_form = "course" in low and any(STUDENT_COLUMNS.get(c) == "section" for c in low)
            for i, c in enumerate(low):
                if long_form and c == "course":
                    columns[i] = ("field", "course")
                elif long_form and STUDENT_COLUMNS.get(c) == "section":
                    columns[i] = ("field", "courseSection")
                elif c in STUDENT_COLUMNS:
                    columns.setdefault(i, ("field", STUDENT_COLUMNS[c]))
                elif norm_name(c) in by_name:
                    columns[i] = ("course", by_name[norm_name(c)])
            continue
        if columns is None:
            continue
        rec, taken = {}, {}
        for i, (kind, key) in columns.items():
            v = cells[i] if i < len(cells) else ""
            if kind == "field":
                rec[key] = v
            elif clean(v):
                taken[key] = clean(v).upper()
        roll = clean_roll(rec.get("roll", ""))
        name = clean_person(rec.get("name", ""))
        if not roll and not name:
            continue
        if not re.fullmatch(r"[A-Z0-9]{6,12}", roll) or not name:
            continue  # sub-headings and totals
        if roll in conflicted:
            continue
        if long_form:
            code = clean(rec.get("course", "")).upper()
            sec = clean(rec.get("courseSection", "")).upper()
            abb = by_code.get(norm_code(code))
            if not code or not re.fullmatch(r"[A-Z]", sec):
                problems.append("%s: course %r, section %r could not be read" % (roll, code, sec))
                continue
            if not abb:
                problems.append("%s: course %r is not in this term's course list" % (roll, code))
                continue
            if roll in students and norm_name(students[roll]["name"]) != norm_name(name):
                problems.append("roll %s has two different names in %s" % (roll, path.name))
                continue
            entry = students.setdefault(roll, {"name": name, "courses": {}})
            if entry["courses"].get(abb, sec) != sec:
                problems.append("%s is in two sections of %s" % (roll, abb))
            entry["courses"][abb] = sec
            continue
        if roll in students:
            if norm_name(students[roll]["name"]) == norm_name(name):
                continue  # the same row printed twice
            # one roll number against two different names: we can't know whose
            # it is, so neither is loaded and the file's owner is told
            notes.append("roll %s is given to two different names in %s; left out until the file is corrected"
                         % (roll, path.name))
            conflicted.add(roll)
            del students[roll]
            continue
        total = rec.get("total", "")
        if taken and clean(total) and int(float(total)) != len(taken):
            problems.append("%s: Total says %s but %d course(s) read" % (roll, clean(total), len(taken)))
        entry = {"name": name}
        for key in ("section", "major", "minor", "track"):
            if clean(rec.get(key, "")):
                entry[key] = clean(rec[key]).upper()
        # a section is one letter; a PDF cell can pick up a stray character
        # from the wrapped name beside it ("W A"), so the last letter stands
        if "section" in entry and not re.fullmatch(r"[A-Z]", entry["section"]):
            last = entry["section"].split()[-1]
            if re.fullmatch(r"[A-Z]", last):
                entry["section"] = last
            else:
                problems.append("%s: section %r could not be read" % (roll, entry["section"]))
        if taken:
            entry["courses"] = taken
        if clean(rec.get("elective", "")):
            abb = match_course(rec["elective"], courses or {})
            if not abb:
                raise ParseError("%s: elective %r matches no course in this term's course list" % (roll, clean(rec["elective"])))
            entry["electives"] = [abb]
        students[roll] = entry
    if problems:
        raise ParseError("; ".join(problems))
    if not students:
        raise ParseError("no students found")
    return students, notes
